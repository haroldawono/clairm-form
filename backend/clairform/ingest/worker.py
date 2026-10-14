"""Document reader executed in a separate, resource-limited process.

Protocol: bytes on stdin, extension in argv[1], limits in argv[2] (JSON); a single JSON
object on stdout: either {"ok": true, ...} or {"ok": false, "code": ..., "detail": ...}.
Isolating parsing in a child process means a malformed or hostile file can at worst crash
or time out this process; it never takes the API server down with it.
"""
from __future__ import annotations

import csv
import io
import json
import sys
import zipfile
import re

MAX_ZIP_ENTRIES = 400
MAX_ZIP_UNCOMPRESSED = 120 * 1024 * 1024
MAX_ZIP_RATIO = 120
MAX_ROWS_PER_SHEET = 5000
MAX_COLS = 60


class ReadError(Exception):
    def __init__(self, code: str, detail: str = ""):
        super().__init__(code)
        self.code = code
        self.detail = detail


def _check_zip(data: bytes, required_member: str) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ReadError("DOC-CORRUPT", "archive bureautique illisible")
    infos = z.infolist()
    total = sum(i.file_size for i in infos)
    if len(infos) > MAX_ZIP_ENTRIES or total > MAX_ZIP_UNCOMPRESSED:
        raise ReadError("DOC-UNSAFE", "archive trop volumineuse une fois décompressée")
    for i in infos:
        name = i.filename.replace("\\", "/")
        if name.startswith("/") or ".." in name.split("/"):
            raise ReadError("DOC-UNSAFE", "chemin interne suspect")
        if i.compress_size and i.file_size / max(1, i.compress_size) > MAX_ZIP_RATIO and i.file_size > 5 * 1024 * 1024:
            raise ReadError("DOC-UNSAFE", "taux de compression anormal")
    if required_member not in z.namelist():
        raise ReadError("DOC-CORRUPT", "structure du document invalide")
    return z


def _clean(text: str, limit: int) -> str:
    # Some PDFs yield lone UTF-16 surrogates; they cannot be encoded and would break the output.
    text = text.encode("utf-8", "replace").decode("utf-8")
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return text[:limit]


def read_pdf(data: bytes, limits: dict) -> dict:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ReadError("DOC-ENCRYPTED", "PDF protégé par mot de passe")
            except ReadError:
                raise
            except Exception:
                raise ReadError("DOC-ENCRYPTED", "PDF protégé par mot de passe")
        count = len(reader.pages)
    except ReadError:
        raise
    except (PdfReadError, Exception) as exc:  # pypdf raises many types on corrupt input
        raise ReadError("DOC-CORRUPT", f"PDF illisible ({exc.__class__.__name__})")
    if count == 0:
        raise ReadError("DOC-CORRUPT", "PDF sans page")
    if count > limits["max_pdf_pages"]:
        raise ReadError("DOC-TOO-MANY-PAGES", f"{count} pages (maximum {limits['max_pdf_pages']})")
    pages = []
    budget = limits["max_text_chars"]
    for n, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = _clean(text, max(0, budget))
        budget -= len(text)
        pages.append({"n": n, "text": text})
    return {"pages": pages, "page_count": count, "tables": [], "meta": {}}


def read_docx(data: bytes, limits: dict) -> dict:
    z = _check_zip(data, "word/document.xml")
    page_count = None
    if "docProps/app.xml" in z.namelist():
        m = re.search(rb"<Pages>(\d+)</Pages>", z.read("docProps/app.xml")[:200000])
        if m:
            page_count = int(m.group(1))
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ReadError("DOC-CORRUPT", f"document Word illisible ({exc.__class__.__name__})")
    lines = [p.text for p in doc.paragraphs]
    tables = []
    for table in doc.tables[:50]:
        rows = []
        for row in table.rows[:MAX_ROWS_PER_SHEET]:
            cells = [c.text.strip() for c in row.cells[:MAX_COLS]]
            rows.append(cells)
            lines.append(" | ".join(c for c in cells if c))
        tables.append({"name": f"Tableau {len(tables) + 1}", "rows": rows})
    text = _clean("\n".join(lines), limits["max_text_chars"])
    return {"pages": [{"n": 1, "text": text}], "page_count": page_count, "tables": tables, "meta": {"page_count_source": "docx-metadata" if page_count else None}}


def read_xlsx(data: bytes, limits: dict) -> dict:
    _check_zip(data, "xl/workbook.xml")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        wb_formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    except Exception as exc:
        raise ReadError("DOC-CORRUPT", f"classeur Excel illisible ({exc.__class__.__name__})")
    tables, lines, uncached = [], [], 0
    for ws, wsf in zip(wb.worksheets[:20], wb_formulas.worksheets[:20]):
        rows = []
        lines.append(f"[Feuille] {ws.title}")
        for values, formulas in zip(ws.iter_rows(values_only=True, max_row=MAX_ROWS_PER_SHEET),
                                    wsf.iter_rows(values_only=True, max_row=MAX_ROWS_PER_SHEET)):
            cells = []
            for v, f in zip(values[:MAX_COLS], formulas[:MAX_COLS]):
                if v is None and isinstance(f, str) and f.startswith("="):
                    uncached += 1
                if v is None:
                    cells.append("")
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    cells.append(v)
                else:
                    cells.append(str(v))
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                rows.append(cells)
                lines.append(" | ".join(str(c) for c in cells if c != ""))
        tables.append({"name": ws.title, "rows": rows})
    wb.close()
    wb_formulas.close()
    text = _clean("\n".join(lines), limits["max_text_chars"])
    return {"pages": [{"n": 1, "text": text}], "page_count": None, "tables": tables, "meta": {"uncached_formulas": uncached}}


def _decode_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ReadError("DOC-CORRUPT", "fichier binaire présenté comme texte")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ReadError("DOC-CORRUPT", "encodage de texte non reconnu")


def read_text(data: bytes, limits: dict) -> dict:
    text = _clean(_decode_text(data), limits["max_text_chars"])
    return {"pages": [{"n": 1, "text": text}], "page_count": None, "tables": [], "meta": {}}


def read_csv(data: bytes, limits: dict) -> dict:
    text = _clean(_decode_text(data), limits["max_text_chars"])
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    if sample.count("\t") > max(sample.count(";"), sample.count(",")):
        delimiter = "\t"
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [c.strip() for c in row[:MAX_COLS]]
        if any(cells):
            rows.append(cells)
        if len(rows) >= MAX_ROWS_PER_SHEET:
            break
    lines = [" | ".join(c for c in r if c) for r in rows]
    return {"pages": [{"n": 1, "text": "\n".join(lines)}], "page_count": None, "tables": [{"name": "CSV", "rows": rows}], "meta": {}}


def read_image(data: bytes, limits: dict) -> dict:
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = 60_000_000
    try:
        with Image.open(io.BytesIO(data)) as im:
            im.verify()
        with Image.open(io.BytesIO(data)) as im:
            width, height = im.size
    except Exception as exc:
        raise ReadError("DOC-CORRUPT", f"image illisible ({exc.__class__.__name__})")
    return {"pages": [{"n": 1, "text": ""}], "page_count": 1, "tables": [], "meta": {"width": width, "height": height}}


READERS = {"pdf": read_pdf, "docx": read_docx, "xlsx": read_xlsx, "txt": read_text, "md": read_text,
           "csv": read_csv, "png": read_image, "jpg": read_image, "jpeg": read_image}


def main() -> None:
    ext = sys.argv[1]
    limits = json.loads(sys.argv[2])
    data = sys.stdin.buffer.read()
    try:
        result = READERS[ext](data, limits)
        result["ok"] = True
    except ReadError as err:
        result = {"ok": False, "code": err.code, "detail": err.detail}
    except ImportError as exc:  # missing library in this interpreter: not the file's fault
        result = {"ok": False, "code": "READER-SETUP", "detail": f"{exc.__class__.__name__}: {exc}"}
    except MemoryError:
        result = {"ok": False, "code": "DOC-UNSAFE", "detail": "mémoire insuffisante pour lire ce fichier"}
    except Exception as exc:  # never leak a traceback to the parent
        result = {"ok": False, "code": "DOC-CORRUPT", "detail": f"lecture impossible ({exc.__class__.__name__})"}
    sys.stdout.write(json.dumps(result, ensure_ascii=True, default=str))
    sys.stdout.flush()


if __name__ == "__main__":
    main()
