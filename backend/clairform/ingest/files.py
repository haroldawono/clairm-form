"""Upload validation and isolated text extraction."""
from __future__ import annotations

import hashlib
import logging
import json
import os
import re
import subprocess
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from ..config import ALLOWED_EXTENSIONS, Settings

BACKEND_DIR = Path(__file__).resolve().parents[2]


class UploadRejected(Exception):
    """A file that cannot be accepted. `code` is stable and documented for the UI."""

    def __init__(self, code: str, message: str, status: int = 400):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status = status


ERROR_MESSAGES = {
    "DOC-EMPTY": "Le fichier est vide.",
    "DOC-FORMAT": "Format non pris en charge.",
    "DOC-MISMATCH": "Le contenu du fichier ne correspond pas à son extension.",
    "DOC-TOO-LARGE": "Le fichier dépasse la taille maximale autorisée.",
    "DOC-CORRUPT": "Le fichier est endommagé ou illisible.",
    "DOC-ENCRYPTED": "Le PDF est protégé par un mot de passe.",
    "DOC-UNSAFE": "Le fichier a une structure anormale et a été refusé par précaution.",
    "DOC-TOO-MANY-PAGES": "Le document contient trop de pages.",
    "DOC-TIMEOUT": "La lecture du fichier a pris trop de temps et a été interrompue.",
    "DOC-NO-TEXT": "Ce document ne contient pas de texte lisible (document scanné ?).",
}


def safe_display_name(name: str | None) -> str:
    """Keep a human-readable name for display only. Files are never stored under this name."""
    name = unicodedata.normalize("NFC", (name or "document").replace("\\", "/").split("/")[-1])
    name = re.sub(r"[\x00-\x1f\x7f<>:\"|?*]+", "", name).strip(" .")
    if len(name) > 140:
        stem, dot, ext = name.rpartition(".")
        name = (stem[:130] + "." + ext) if dot and len(ext) <= 5 else name[:140]
    return name or "document"


def extension_of(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def check_signature(ext: str, data: bytes) -> None:
    if not data:
        raise UploadRejected("DOC-EMPTY", ERROR_MESSAGES["DOC-EMPTY"])
    if ext not in ALLOWED_EXTENSIONS:
        raise UploadRejected("DOC-FORMAT", ERROR_MESSAGES["DOC-FORMAT"])
    head = data[:1024]
    ok = {
        "pdf": b"%PDF-" in head,
        "docx": data.startswith(b"PK\x03\x04"),
        "xlsx": data.startswith(b"PK\x03\x04"),
        "png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        "jpg": data.startswith(b"\xff\xd8\xff"),
        "jpeg": data.startswith(b"\xff\xd8\xff"),
    }.get(ext, b"\x00" not in head)
    if not ok:
        raise UploadRejected("DOC-MISMATCH", ERROR_MESSAGES["DOC-MISMATCH"])


@dataclass
class Extracted:
    sha256: str
    ext: str
    pages: list[dict]
    page_count: int | None
    tables: list[dict]
    meta: dict

    @property
    def text(self) -> str:
        return "\n\n".join(p["text"] for p in self.pages)

    @property
    def has_text(self) -> bool:
        return len(re.sub(r"\s+", "", self.text)) >= 20


def _limit_resources():  # pragma: no cover - runs in the child process
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
        resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))  # the reader never writes files
    except Exception:
        pass


log = logging.getLogger("clairform.ingest")
_MODE = {"isolated": True, "reason": ""}
from collections import OrderedDict  # noqa: E402

_CACHE: "OrderedDict[tuple, Extracted]" = OrderedDict()  # last files read, keyed by content hash


def _child_env() -> dict:
    env = {k: v for k, v in os.environ.items() if k not in {"PYTHONSTARTUP", "PYTHONINSPECT", "PYTHONBREAKPOINT"}}
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def _read_isolated(ext: str, data: bytes, limits: dict, timeout: int) -> dict | None:
    """Run the reader in a child process. Returns None when the child process itself cannot work
    on this machine (interpreter or library problem), so the caller can fall back to in-process reading."""
    # The child receives the parent's exact sys.path: libraries installed "for the user"
    # (pip install --user, the Windows default without admin rights) must be found too.
    code = (
        "import sys, json; sys.path[:0] = json.loads(sys.argv.pop(1)); "
        "from clairform.ingest.worker import main; main()"
    )
    paths = [str(BACKEND_DIR)] + [p for p in sys.path if p]
    try:
        proc = subprocess.run(
            [sys.executable, "-c", code, json.dumps(paths), ext, json.dumps(limits)],
            input=data, capture_output=True, timeout=timeout, env=_child_env(),
            preexec_fn=_limit_resources if os.name == "posix" else None,
        )
    except subprocess.TimeoutExpired:
        raise UploadRejected("DOC-TIMEOUT", ERROR_MESSAGES["DOC-TIMEOUT"], 422)
    except OSError as exc:
        _MODE.update(isolated=False, reason=f"processus de lecture impossible à lancer ({exc})")
        return None
    try:
        result = json.loads(proc.stdout.decode("utf-8") or "{}")
    except (UnicodeDecodeError, json.JSONDecodeError):
        result = {}
    if not result:
        stderr = proc.stderr.decode("utf-8", "replace")[-800:]
        if "MemoryError" in stderr or proc.returncode in (-9, 137):
            raise UploadRejected("DOC-UNSAFE", ERROR_MESSAGES["DOC-UNSAFE"] + " (mémoire)", 422)
        # A crash on one particular file: read this file in-process, keep isolation for the next ones.
        log.warning("processus de lecture en échec (code %s) : %s", proc.returncode, stderr.strip()[-300:])
        return None
    if result.get("code") == "READER-SETUP":
        _MODE.update(isolated=False, reason=f"bibliothèque introuvable dans le processus de lecture : {result.get('detail')}")
        return None
    return result


def _read_in_process(ext: str, data: bytes, limits: dict, timeout: int) -> dict:
    from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeout

    from . import worker

    def run():
        try:
            return {"ok": True, **worker.READERS[ext](data, limits)}
        except worker.ReadError as err:
            return {"ok": False, "code": err.code, "detail": err.detail}
        except ImportError as exc:
            return {"ok": False, "code": "READER-SETUP", "detail": f"{exc.__class__.__name__}: {exc}"}
        except Exception as exc:
            log.exception("lecture du fichier impossible")
            return {"ok": False, "code": "DOC-CORRUPT", "detail": f"lecture impossible ({exc.__class__.__name__})"}

    pool = ThreadPoolExecutor(max_workers=1)
    try:
        return pool.submit(run).result(timeout=timeout)
    except FutureTimeout:
        raise UploadRejected("DOC-TIMEOUT", ERROR_MESSAGES["DOC-TIMEOUT"], 422)
    finally:
        pool.shutdown(wait=False)


def reader_status() -> dict:
    return {"isolated": _MODE["isolated"], "reason": _MODE["reason"]}


def clear_cache() -> None:
    _CACHE.clear()


def extract(filename: str, data: bytes, settings: Settings) -> Extracted:
    ext = extension_of(filename)
    if len(data) > settings.max_upload_bytes:
        raise UploadRejected("DOC-TOO-LARGE", ERROR_MESSAGES["DOC-TOO-LARGE"], 413)
    check_signature(ext, data)
    limits = {"max_pdf_pages": settings.max_pdf_pages, "max_text_chars": settings.max_text_chars}
    sha = hashlib.sha256(data).hexdigest()
    key = (sha, ext, settings.max_pdf_pages, settings.max_text_chars)
    cached = _CACHE.get(key)
    if cached is not None:  # identical bytes already read (e.g. the demonstration files): no second parse
        _CACHE.move_to_end(key)
        return cached
    result = None
    if _MODE["isolated"] and settings.reader != "inprocess":
        result = _read_isolated(ext, data, limits, settings.extraction_timeout_s)
        if result is None:
            log.warning("Lecture isolée indisponible, passage en lecture dans le processus principal : %s", _MODE["reason"])
    if result is None:
        result = _read_in_process(ext, data, limits, settings.extraction_timeout_s)
    if not result.get("ok"):
        code_ = result.get("code") or "DOC-CORRUPT"
        if code_ == "READER-SETUP":
            raise UploadRejected("READER-SETUP", "Lecture impossible : une bibliothèque nécessaire n'est pas installée sur ce serveur "
                                 f"({result.get('detail')}). Relancez « pip install -r requirements.txt » puis redémarrez ClairForm.", 500)
        message = ERROR_MESSAGES.get(code_, ERROR_MESSAGES["DOC-CORRUPT"])
        detail = result.get("detail")
        raise UploadRejected(code_, f"{message} ({detail})" if detail else message, 422)
    extracted = Extracted(
        sha256=sha,
        ext=ext,
        pages=result["pages"],
        page_count=result.get("page_count"),
        tables=result.get("tables") or [],
        meta=result.get("meta") or {},
    )
    _CACHE[key] = extracted
    while len(_CACHE) > 48:
        _CACHE.popitem(last=False)
    return extracted


def self_test(settings: Settings) -> list[str]:
    """Read one small file of each format at start-up. Returns the list of problems (empty when all is well)."""
    import io

    problems = []
    samples_: dict[str, bytes] = {"txt": "Pièces à fournir : statuts.".encode("utf-8")}
    try:
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "Test de lecture ClairForm")
        c.save()
        samples_["pdf"] = buf.getvalue()
    except Exception as exc:
        problems.append(f"génération PDF de test impossible : {exc}")
    try:
        from docx import Document

        buf = io.BytesIO()
        d = Document()
        d.add_paragraph("Test de lecture ClairForm")
        d.save(buf)
        samples_["docx"] = buf.getvalue()
    except Exception as exc:
        problems.append(f"python-docx : {exc}")
    try:
        from openpyxl import Workbook

        buf = io.BytesIO()
        wb = Workbook()
        wb.active.append(["Total", 1])
        wb.save(buf)
        samples_["xlsx"] = buf.getvalue()
    except Exception as exc:
        problems.append(f"openpyxl : {exc}")
    for ext, data in samples_.items():
        try:
            extract(f"test.{ext}", data, settings)
        except UploadRejected as err:
            problems.append(f"lecture {ext.upper()} : {err.message}")
    return problems
