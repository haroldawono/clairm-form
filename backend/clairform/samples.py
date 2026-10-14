"""Demonstration documents: a call for proposals and the pieces of an applicant's file.

They are laid out like real administrative documents (letterheads, reference numbers, tables,
signature blocks, page numbers). The programme, the association and the people are invented;
documents that would normally be issued by an administration (identity card, registry extract)
carry a "SPÉCIMEN" watermark so they can never be mistaken for genuine papers.
Dates are computed relative to `today` so the scenario stays meaningful whenever it is run.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from .analysis.text import add_months, format_date_fr

PROGRAM_NAME = "Fonds Tremplin Impact 2027"
FUNDER = "Agence Tremplin pour l'Innovation Solidaire"
DEMO_TITLE = "Kombi Solaire — Fonds Tremplin Impact 2027"
ORG = "Association Kombi Solaire"
ORG_ADDRESS = "Rue des Palmiers, quartier Bonamoussadi — BP 4512 Douala"
LEADER = "Amina Ekané"


# ----------------------------------------------------------------------------- PDF toolkit
def _styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    ss = getSampleStyleSheet()
    ink = colors.HexColor("#1d2733")
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=ink, spaceAfter=4),
        "subtitle": ParagraphStyle("st", parent=ss["Normal"], fontName="Helvetica", fontSize=10.5, leading=14, alignment=TA_CENTER,
                                   textColor=colors.HexColor("#4a5565"), spaceAfter=10),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=ink, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=10, leading=14, alignment=TA_JUSTIFY, textColor=ink),
        "small": ParagraphStyle("s", parent=ss["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#4a5565")),
        "cell": ParagraphStyle("c", parent=ss["BodyText"], fontName="Helvetica", fontSize=9, leading=12, textColor=ink),
        "cellb": ParagraphStyle("cb", parent=ss["BodyText"], fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=ink),
        "right": ParagraphStyle("r", parent=ss["BodyText"], fontName="Helvetica", fontSize=10, leading=14, alignment=2, textColor=ink),
    }


def _letterhead_painter(org: str, lines: list[str], color: str, running: str | None = None, ref: str | None = None,
                        specimen: bool = False):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4

    w, h = A4

    def paint(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(color))
        canvas.rect(0, h - 10, w, 10, stroke=0, fill=1)
        if doc.page == 1:
            canvas.setFillColor(colors.HexColor(color))
            canvas.circle(62, h - 52, 15, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont("Helvetica-Bold", 12)
            canvas.drawCentredString(62, h - 56, "".join(word[0] for word in org.split()[:3] if word[0].isupper())[:3])
            canvas.setFillColor(colors.HexColor("#1d2733"))
            canvas.setFont("Helvetica-Bold", 11.5)
            canvas.drawString(86, h - 48, org)
            canvas.setFont("Helvetica", 8.5)
            canvas.setFillColor(colors.HexColor("#4a5565"))
            for i, line in enumerate(lines):
                canvas.drawString(86, h - 61 - i * 10.5, line)
            canvas.setStrokeColor(colors.HexColor("#c9d0d9"))
            canvas.line(48, h - 72 - len(lines) * 10.5, w - 48, h - 72 - len(lines) * 10.5)
        elif running:
            canvas.setFont("Helvetica", 8.5)
            canvas.setFillColor(colors.HexColor("#6b7480"))
            canvas.drawString(48, h - 34, running)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6b7480"))
        if ref:
            canvas.drawString(48, 30, ref)
        canvas.drawRightString(w - 48, 30, f"Page {doc.page}")
        if specimen:
            canvas.setFillColor(colors.Color(0.75, 0.1, 0.1, alpha=0.12))
            canvas.setFont("Helvetica-Bold", 78)
            canvas.translate(w / 2, h / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, "SPÉCIMEN")
        canvas.restoreState()

    return paint


def _pdf(story, painter, top_margin=120) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=top_margin, bottomMargin=60,
                            title=story[0].text if hasattr(story[0], "text") else "Document", author="")
    doc.build(story, onFirstPage=painter, onLaterPages=painter)
    return out.getvalue()


def _table(rows, widths, header=True, zebra=False):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    st = _styles()
    data = [[Paragraph(str(c), st["cellb"] if (header and i == 0) else st["cell"]) for c in row] for i, row in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9c1cc")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9edf2")))
    t.setStyle(TableStyle(style))
    return t


class _Stamp:
    """A round rubber-stamp drawing (flowable) placed next to a signature."""

    def __new__(cls, org: str, city: str = "DOUALA", color: str = "#2445c7"):
        from reportlab.graphics.shapes import Circle, Drawing, String
        from reportlab.lib import colors

        d = Drawing(120, 80)
        c = colors.HexColor(color)
        d.add(Circle(60, 40, 36, strokeColor=c, fillColor=None, strokeWidth=2))
        d.add(Circle(60, 40, 29, strokeColor=c, fillColor=None, strokeWidth=0.8))
        d.add(String(60, 44, org.upper()[:14], fontName="Helvetica-Bold", fontSize=6.5, fillColor=c, textAnchor="middle"))
        d.add(String(60, 33, city, fontName="Helvetica", fontSize=7, fillColor=c, textAnchor="middle"))
        return d


def _signature_block(place_date: str, role: str, name: str, stamp_org: str | None = None):
    from reportlab.platypus import Paragraph, Spacer, Table

    st = _styles()
    left = [Paragraph(place_date, st["body"]), Spacer(1, 6), Paragraph(f"<b>{role}</b>", st["body"]), Paragraph(name, st["body"]),
            Spacer(1, 4), Paragraph("<i>Signature</i> : (signé)", st["body"])]
    right = _Stamp(stamp_org) if stamp_org else ""
    t = Table([[left, right]], colWidths=[300, 140])
    t.setStyle([("VALIGN", (0, 0), (-1, -1), "TOP")])
    return t


# ----------------------------------------------------------------------------- the call
def deadline_for(today: date) -> date:
    return today + timedelta(days=40)


def call_pdf(today: date) -> bytes:
    from reportlab.platypus import PageBreak, Paragraph, Spacer

    st = _styles()
    deadline = deadline_for(today)
    P = lambda t, s="body": Paragraph(t, st[s])  # noqa: E731
    story = [
        P("Appel à projets — Fonds Tremplin Impact 2027", "title"),
        P(f"Règlement de l'appel à projets · Édition 2027 · Publié le {format_date_fr(today - timedelta(days=12))}", "subtitle"),
        P("Article 1 – Objet de l'appel", "h2"),
        P("Le Fonds Tremplin Impact soutient des projets à impact social ou environnemental portés par de jeunes structures. "
          "Il finance la phase de lancement d'une solution déjà testée auprès de ses premiers bénéficiaires, dans les domaines de "
          "l'énergie, de l'agriculture, de l'eau et de l'assainissement, de l'éducation et de la santé communautaire."),
        P("Le présent règlement précise les conditions de participation, la composition du dossier et les modalités de sélection. "
          "Il est accompagné d'un formulaire de candidature et d'un modèle de budget, téléchargeables sur la plateforme du Fonds."),
        P("Article 2 – Éligibilité", "h2"),
        P("Peuvent candidater les structures répondant à toutes les conditions suivantes :"),
        P("- Être une association ou une entreprise légalement constituée depuis au moins 12 mois ;"),
        P("- Avoir son siège social dans un pays d'Afrique centrale ou de l'Ouest ;"),
        P("- Le porteur de projet doit être âgé de 18 à 35 ans à la date limite de dépôt."),
        P("Article 3 – Dépenses prises en charge", "h2"),
        P("Les dépenses prises en charge comprennent notamment les frais de personnel affecté au projet, l'achat d'équipements, "
          "la formation des bénéficiaires, la communication et une part des frais de fonctionnement. Les dépenses d'investissement "
          "immobilier et le remboursement de dettes antérieures ne sont pas financés."),
        P("Article 4 – Financement", "h2"),
        P("Le montant de la subvention ne peut excéder 50 000 € par projet."),
        P("Le candidat doit justifier d'un cofinancement (apport propre ou autre financeur) d'au moins 20 % du budget total."),
        P("La durée du projet ne doit pas excéder 12 mois."),
        P("La subvention est versée en deux tranches : 70 % à la signature de la convention et 30 % à la remise du rapport final."),
        PageBreak(),
        P("Article 5 – Pièces à fournir", "h2"),
        P("Le dossier de candidature doit obligatoirement comprendre les pièces suivantes, sous peine d'irrecevabilité :"),
        P("- Le formulaire de candidature complété et signé ;"),
        P("- Une copie de la pièce d'identité du porteur de projet en cours de validité ;"),
        P("- Un extrait d'immatriculation au registre (RCCM ou récépissé de déclaration) datant de moins de 3 mois ;"),
        P("- Une note de présentation du projet (problème, solution, bénéficiaires, impact attendu) ;"),
        P("- Le budget prévisionnel détaillé faisant apparaître le montant demandé et l'apport propre ;"),
        P("- Un calendrier de mise en œuvre (chronogramme) ;"),
        P("- Deux CV : celui du porteur et celui d'un second membre clé de l'équipe ;"),
        P("- Au moins une lettre de soutien d'un partenaire, signée ;"),
        P("- Une attestation sur l'honneur signée par le représentant légal ;"),
        P("- Facultatif : des photos ou une courte vidéo de présentation du projet."),
        P("La note de présentation ne doit pas dépasser 10 pages."),
        P("Tous les documents doivent être transmis au format PDF, à l'exception du budget qui peut être fourni au format Excel."),
        P("Article 6 – Calendrier et modalités de dépôt", "h2"),
        P(f"Date limite de dépôt des candidatures : {format_date_fr(deadline)} à 23h59 (heure de Douala)."),
        P("Les dossiers sont déposés en ligne sur la plateforme du Fonds. Aucun dossier papier n'est accepté. "
          "Un accusé de réception est envoyé par courriel dans les 48 heures."),
        P("Article 7 – Critères de sélection", "h2"),
        P("Les projets seront évalués selon leur pertinence, leur faisabilité et leur impact. Le comité de sélection "
          "s'engage à notifier sa décision à chaque candidat dans un délai de huit semaines après la clôture."),
        P("Article 8 – Contact", "h2"),
        P("Pour toute question : Secrétariat du Fonds Tremplin Impact, appels@tremplin-impact.org, du lundi au vendredi de 9h à 16h."),
    ]
    painter = _letterhead_painter(FUNDER, ["Fonds Tremplin Impact — Secrétariat technique", "BP 2210 Douala · appels@tremplin-impact.org"],
                                  "#0f5c4d", running="Fonds Tremplin Impact 2027 — Règlement de l'appel à projets",
                                  ref="Réf. FTI/AAP/2027-01")
    return _pdf(story, painter)


# ----------------------------------------------------------------------------- applicant pieces
def _org_painter(ref: str | None = None):
    return _letterhead_painter(ORG, [ORG_ADDRESS, "Tél. +237 6 99 00 00 00 · contact@kombisolaire.org · Récépissé n° 0457/RDA/J06/BAPP"],
                               "#c27c0e", ref=ref)


def application_form(today: date, requested: int = 45000) -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    st = _styles()
    P = lambda t, s="body": Paragraph(t, st[s])  # noqa: E731
    story = [
        P("Formulaire de candidature — Fonds Tremplin Impact 2027", "title"),
        P("Dossier n° FTI-2027-0318 · À remplir par le porteur de projet", "subtitle"),
        P("1. Identification de la structure", "h2"),
        _table([["Rubrique", "Renseignement"], ["Nom de la structure", "Kombi Solaire"], ["Forme juridique", "Association (loi n° 90/053)"],
                ["Numéro d'enregistrement", "0457/RDA/J06/BAPP"], ["Date de création", "14 mars 2023"], ["Siège social", ORG_ADDRESS],
                ["Courriel", "contact@kombisolaire.org"]], [170, 270]),
        P("2. Représentant légal et porteur de projet", "h2"),
        _table([["Rubrique", "Renseignement"], ["Nom et prénom", LEADER], ["Fonction", "Présidente et coordinatrice du projet"],
                ["Date de naissance", "14/03/1996"], ["Coordonnées du porteur", "+237 6 99 00 00 01 · a.ekane@kombisolaire.org"]], [170, 270]),
        P("3. Le projet", "h2"),
        _table([["Rubrique", "Renseignement"], ["Intitulé du projet", "Kits solaires partagés pour les marchés de quartier"],
                ["Lieu de mise en œuvre", "Douala (marchés de Bonamoussadi, Ndokoti et Mboppi)"], ["Durée", "12 mois"],
                ["Bénéficiaires directs", "120 commerçantes et commerçants"]], [170, 270]),
        P("4. Financement", "h2"),
        _table([["Rubrique", "Montant"], ["Coût total du projet", "61 500 €"], ["Montant demandé", f"{requested:,} €".replace(",", " ")],
                ["Apport propre de l'association", "12 500 €"], ["Autres financements", "4 000 €"]], [270, 170]),
        Spacer(1, 10),
        P("Je certifie l'exactitude des informations fournies dans le présent formulaire et m'engage à informer le Fonds de toute modification."),
        Spacer(1, 10),
        _signature_block(f"Fait à Douala, le {format_date_fr(today - timedelta(days=2))}", "La représentante légale", LEADER, "Kombi Solaire"),
    ]
    return _pdf(story, _org_painter("Formulaire FTI-2027 · v1"))


def identity_card(today: date) -> bytes:
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table

    st = _styles()
    P = lambda t, s="body": Paragraph(t, st[s])  # noqa: E731
    card = Table([
        [P("<b>CARTE NATIONALE D'IDENTITÉ</b>"), ""],
        [P("[photo]", "small"), P(f"Nom : <b>EKANÉ</b><br/>Prénoms : <b>Amina Grâce</b><br/>Date de naissance : 14/03/1996<br/>"
                                  f"Lieu de naissance : Douala<br/>Sexe : F · Taille : 1,68 m<br/>Profession : Ingénieure")],
        [P(f"Date de délivrance : {(today - timedelta(days=700)).strftime('%d/%m/%Y')}<br/>"
           f"Date d'expiration : {(today + timedelta(days=1500)).strftime('%d/%m/%Y')}<br/>N° 1000 2345 6789", "small"), ""],
    ], colWidths=[110, 300])
    card.setStyle([("BOX", (0, 0), (-1, -1), 1.2, colors.HexColor("#0f5c4d")), ("SPAN", (0, 0), (1, 0)), ("SPAN", (0, 2), (1, 2)),
                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e3f2ea")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("BOX", (0, 1), (0, 1), 0.6, colors.HexColor("#9aa6b5")), ("TOPPADDING", (0, 0), (-1, -1), 6)])
    story = [P("Copie de la pièce d'identité du porteur de projet", "title"),
             P("Recto de la carte — copie jointe au dossier de candidature", "subtitle"), Spacer(1, 12), card, Spacer(1, 18),
             P("Copie certifiée conforme à l'original.", "small")]
    painter = _letterhead_painter("Pièce d'identité", ["Copie jointe au dossier FTI-2027-0318"], "#0f5c4d", specimen=True)
    return _pdf(story, painter)


def registry_extract(today: date, issued: date) -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    st = _styles()
    P = lambda t, s="body": Paragraph(t, st[s])  # noqa: E731
    story = [
        P("Extrait d'immatriculation au registre des associations", "title"),
        P("Registre du commerce et du crédit mobilier — extrait délivré sur demande", "subtitle"),
        _table([["Rubrique", "Inscription"], ["Dénomination", "Kombi Solaire"], ["Forme", "Association déclarée"],
                ["Numéro d'enregistrement", "0457/RDA/J06/BAPP"], ["Date de déclaration", "14/03/2023"], ["Siège social", ORG_ADDRESS],
                ["Objet social", "Faciliter l'accès des petits commerces à une énergie propre et abordable"],
                ["Représentante légale", f"{LEADER}, présidente"], ["Situation", "Active, aucune mention de dissolution"]], [170, 270]),
        Spacer(1, 12),
        P(f"Délivré le {format_date_fr(issued)}, pour servir et valoir ce que de droit."),
        Spacer(1, 18),
        P("Le greffier du registre", "right"),
    ]
    painter = _letterhead_painter("Registre des associations", ["Service du registre — Douala", "Extrait n° RA-DLA-2023-0457"],
                                  "#5b6573", ref="Extrait RA-DLA-2023-0457", specimen=True)
    return _pdf(story, painter)


NOTE_SECTIONS = [
    ("Résumé du projet",
     "Kombi Solaire propose des kits solaires partagés, loués à la soirée, aux commerçantes et commerçants des marchés de quartier de "
     "Douala. Une station de recharge commune, gérée par une coopérative de commerçants, alimente des lampes et des batteries portatives. "
     "Le projet vise à allonger les heures de vente après la tombée de la nuit et à remplacer les lampes à pétrole."),
    ("Contexte et problème",
     "Dans les marchés de Bonamoussadi, Ndokoti et Mboppi, une grande partie des étals ne disposent d'aucun éclairage fiable. Les "
     "coupures du réseau et le coût d'un raccordement individuel obligent les vendeuses à fermer tôt ou à s'éclairer au pétrole, avec "
     "des risques d'incendie et une fumée nocive. Nos enquêtes auprès de 64 commerçantes en 2025 montrent une perte de chiffre "
     "d'affaires en fin de journée et une dépense récurrente en combustible."),
    ("Solution proposée",
     "Chaque marché accueille une station solaire de recharge et un parc de 60 kits (lampe LED, batterie, prise USB). Les kits sont loués "
     "à la soirée pour un montant inférieur à la dépense actuelle en pétrole. La coopérative encaisse les locations, assure la petite "
     "maintenance et reverse une part à un fonds de renouvellement du matériel."),
    ("Bénéficiaires",
     "Les bénéficiaires directs sont 120 commerçantes et commerçants, dont une majorité de femmes vendeuses de produits frais. "
     "Les bénéficiaires indirects sont leurs familles et leurs clients, qui profitent de marchés éclairés et plus sûrs le soir."),
    ("Impact attendu",
     "À douze mois, nous visons une heure et demie de vente supplémentaire par soirée, la suppression des lampes à pétrole sur les étals "
     "équipés et la formation de six techniciennes et techniciens de maintenance issus des marchés."),
    ("Modèle économique",
     "Les recettes de location couvrent la maintenance, le salaire des gestionnaires de station et le renouvellement des batteries. "
     "La subvention finance l'investissement initial et la formation ; l'apport propre de l'association couvre la coordination."),
    ("Équipe",
     "Amina Ekané, ingénieure en énergies renouvelables, coordonne le projet. Paul Mbarga, technicien photovoltaïque, dirige l'installation "
     "et la formation des techniciens. Un comité de trois représentantes des commerçantes suit la mise en œuvre."),
    ("Partenaires",
     "La coopérative des commerçants du marché central met à disposition un espace de démonstration et gère les locations. "
     "La mairie d'arrondissement facilite l'implantation des stations."),
    ("Calendrier de mise en œuvre",
     "Mois 1 à 3 : achat et assemblage des kits. Mois 3 à 6 : installation des stations. Mois 6 à 9 : formation et lancement des "
     "locations. Mois 9 à 12 : évaluation, ajustements et bilan."),
    ("Budget résumé",
     "Le coût total s'élève à 61 500 €. Le montant demandé au Fonds est de 45 000 €, complété par un apport propre de 12 500 € et un "
     "financement municipal de 4 000 €."),
    ("Risques et mesures",
     "Vol de matériel : stockage nocturne dans la station fermée. Panne de batteries : stock tampon de 10 % et garantie fournisseur. "
     "Faible adhésion : tarifs testés lors d'un pilote de trois mois en 2025."),
    ("Suivi et évaluation",
     "Indicateurs suivis chaque mois : nombre de locations, heures de vente supplémentaires déclarées, incidents techniques, recettes. "
     "Un rapport intermédiaire est remis au sixième mois et un rapport final au douzième mois."),
]


def project_note(today: date, pages: int) -> bytes:
    from reportlab.platypus import PageBreak, Paragraph

    st = _styles()
    story = [Paragraph("Note de présentation du projet", st["title"]),
             Paragraph("Kits solaires partagés pour les marchés de quartier — Fonds Tremplin Impact 2027", st["subtitle"])]
    for i in range(pages):
        title, text = NOTE_SECTIONS[i % len(NOTE_SECTIONS)]
        if i:
            story.append(PageBreak())
        story.append(Paragraph(f"{i + 1}. {title}", st["h2"]))
        story.append(Paragraph(text, st["body"]))
        story.append(Paragraph(text, st["body"]) if len(text) < 400 else Paragraph("", st["body"]))
    return _pdf(story, _org_painter("Note de présentation · Kombi Solaire"))


def cv(name: str, title: str, experience: list[tuple[str, str]], education: list[str], skills: str) -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    st = _styles()
    story = [Paragraph(f"Curriculum vitae — {name}", st["title"]), Paragraph(title, st["subtitle"]),
             Paragraph("Expérience professionnelle", st["h2"])]
    for period, text in experience:
        story.append(Paragraph(f"<b>{period}</b> — {text}", st["body"]))
        story.append(Spacer(1, 4))
    story.append(Paragraph("Formation", st["h2"]))
    for e in education:
        story.append(Paragraph(e, st["body"]))
    story.append(Paragraph("Compétences", st["h2"]))
    story.append(Paragraph(skills, st["body"]))
    story.append(Paragraph("Langues", st["h2"]))
    story.append(Paragraph("Français (langue maternelle), anglais (courant), duala (courant).", st["body"]))
    painter = _letterhead_painter(name, ["Douala, Cameroun", "Membre de l'équipe Kombi Solaire"], "#2445c7")
    return _pdf(story, painter)


def declaration_docx(today: date) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = Document()
    p = d.add_paragraph(ORG)
    p.runs[0].bold = True
    d.add_paragraph(ORG_ADDRESS)
    h = d.add_heading("Attestation sur l'honneur", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    body = (f"Je soussignée {LEADER}, agissant en qualité de représentante légale de l'association Kombi Solaire, dont le siège "
            "est situé à Douala, atteste sur l'honneur :")
    d.add_paragraph(body)
    for item in ["l'exactitude des informations fournies dans le dossier de candidature au Fonds Tremplin Impact 2027 ;",
                 "que l'association est à jour de ses obligations déclaratives ;",
                 "que le projet présenté ne bénéficie pas d'un autre financement pour les mêmes dépenses ;",
                 "que l'association s'engage à utiliser la subvention conformément à la convention qui sera signée."]:
        d.add_paragraph(item, style="List Bullet")
    d.add_paragraph(f"Fait à Douala, le {format_date_fr(today - timedelta(days=1))}.")
    sig = d.add_paragraph("La représentante légale\n" + LEADER + "\nSignature : (signé)")
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for para in d.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)
    props = d.core_properties
    props.author = LEADER
    props.title = "Attestation sur l'honneur"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def support_letter(today: date) -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    st = _styles()
    P = lambda t, s="body": Paragraph(t, st[s])  # noqa: E731
    story = [
        P(f"Douala, le {format_date_fr(today - timedelta(days=3))}", "right"),
        Spacer(1, 6),
        P("À l'attention du comité de sélection<br/>Fonds Tremplin Impact 2027"),
        Spacer(1, 10),
        P("<b>Objet : lettre de soutien au projet « Kits solaires partagés pour les marchés de quartier »</b>"),
        Spacer(1, 8),
        P("Madame, Monsieur,"),
        P("La Coopérative des commerçants du marché central apporte son soutien au projet porté par l'association Kombi Solaire. "
          "Nos membres subissent chaque soir l'absence d'éclairage fiable sur les étals, et nous avons participé au pilote de trois mois "
          "mené en 2025."),
        P("En qualité de partenaire, nous nous engageons à mettre à disposition un espace de démonstration, à assurer la gestion des "
          "locations de kits auprès de nos membres et à désigner trois représentantes au comité de suivi du projet."),
        P("Nous vous prions d'agréer, Madame, Monsieur, l'expression de nos salutations distinguées."),
        Spacer(1, 14),
        _signature_block("", "Le président", "Joseph Ndoumbé", "Coopérative"),
    ]
    painter = _letterhead_painter("Coopérative des commerçants du marché central",
                                  ["Marché central, porte 3 — BP 1187 Douala", "coop.marchecentral@gmail.com"], "#7a3e9d",
                                  ref="Réf. CCMC/2026/112")
    return _pdf(story, painter)


# ----------------------------------------------------------------------------- spreadsheets
def _xlsx_budget(corrected: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget prévisionnel"
    thin = Side(style="thin", color="B9C1CC")
    head_fill = PatternFill("solid", fgColor="E9EDF2")
    total_fill = PatternFill("solid", fgColor="FFF3D6")
    rows = [
        ["Budget prévisionnel détaillé — Kits solaires partagés pour les marchés de quartier"],
        ["Structure : Association Kombi Solaire", "", "", "Devise : EUR"],
        [],
        ["Poste de dépense", "Quantité", "Coût unitaire (EUR)", "Montant (EUR)"],
        ["Personnel (coordination et techniciens, 12 mois)", 12, 2000, 24000],
        ["Équipement (kits LED, batteries, 3 stations)", 1, 18500, 18500],
        ["Formation des commerçants et des techniciens", 6, 1500, 9000],
        ["Communication et sensibilisation", 1, 4000, 4000],
        ["Frais de fonctionnement", 12, 500, 6000],
        ["Total dépenses", "", "", 61500 if corrected else 62000],
        [],
        ["Ressources", "", "", "Montant (EUR)"],
        ["Montant demandé au Fonds Tremplin", "", "", 45000],
        ["Apport propre de l'association", "", "", 12500 if corrected else 13000],
        ["Autre financeur (mairie d'arrondissement)", "", "", 4000],
        ["Total ressources", "", "", 61500 if corrected else 62000],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        label = str(row[0].value or "")
        for cell in row:
            if cell.value not in (None, ""):
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
        if label in ("Poste de dépense", "Ressources"):
            for cell in row:
                cell.fill = head_fill
                cell.font = Font(bold=True)
        if label.startswith("Total"):
            for cell in row:
                cell.fill = total_fill
                cell.font = Font(bold=True)
    for col, width in zip("ABCD", (52, 11, 20, 18)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def budget(today: date, corrected: bool) -> bytes:
    return _xlsx_budget(corrected)


def timeline(today: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    start = date(today.year + 1, 2, 1)
    wb = Workbook()
    ws = wb.active
    ws.title = "Chronogramme"
    ws.append(["Chronogramme de mise en œuvre — Kits solaires partagés"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    header = ["Activité", "Responsable", "Début", "Fin"] + [f"M{i}" for i in range(1, 13)]
    ws.append(header)
    acts = [("Achat et assemblage des kits", "Paul Mbarga", 0, 2), ("Installation des stations", "Paul Mbarga", 2, 5),
            ("Formation et lancement des locations", LEADER, 5, 8), ("Évaluation et bilan", LEADER, 8, 11)]
    fill = PatternFill("solid", fgColor="2445C7")
    for name, who, m0, m1 in acts:
        begin = add_months(start, m0)
        end = add_months(start, m1) + (timedelta(days=14) if m1 == 11 else timedelta(days=0))
        ws.append([name, who, datetime(begin.year, begin.month, begin.day), datetime(end.year, end.month, end.day)] +
                  ["" for _ in range(12)])
        r = ws.max_row
        ws.cell(r, 3).number_format = "DD/MM/YYYY"
        ws.cell(r, 4).number_format = "DD/MM/YYYY"
        for m in range(m0, max(m0 + 1, m1)):
            ws.cell(r, 5 + m).fill = fill
    for row in ws.iter_rows(min_row=3, max_row=3):
        for c in row:
            c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- scenario
def demo_pieces(today: date) -> list[tuple[str, bytes]]:
    return [
        ("formulaire_candidature.pdf", application_form(today)),
        ("CNI_porteur.pdf", identity_card(today)),
        ("extrait_RCCM.pdf", registry_extract(today, add_months(today, -5))),
        ("note_presentation_projet.pdf", project_note(today, pages=12)),
        ("budget_previsionnel.xlsx", budget(today, corrected=False)),
        ("chronogramme.xlsx", timeline(today)),
        ("CV_Amina_Ekane.pdf", cv("Amina Ekané", "Ingénieure en énergies renouvelables — coordinatrice du projet", [
            ("2021-2025", "Cheffe de projet énergie solaire communautaire, ONG Lumière des Quartiers, Douala : 14 mini-réseaux installés."),
            ("2019-2021", "Ingénieure d'études, bureau d'études énergie : dimensionnement d'installations photovoltaïques."),
            ("2018", "Stage de fin d'études : audit énergétique de marchés municipaux."),
        ], ["Master en ingénierie des énergies renouvelables (2019)", "Licence de physique appliquée (2017)"],
            "Gestion de projet, dimensionnement solaire, animation communautaire, suivi-évaluation.")),
        ("attestation_sur_honneur.docx", declaration_docx(today)),
    ]


def correction_pieces(today: date) -> list[tuple[str, bytes]]:
    return [
        ("budget_previsionnel_v2.xlsx", budget(today, corrected=True)),
        ("extrait_RCCM_recent.pdf", registry_extract(today, today - timedelta(days=6))),
        ("note_presentation_projet_v2.pdf", project_note(today, pages=8)),
        ("CV_Paul_Mbarga.pdf", cv("Paul Mbarga", "Technicien photovoltaïque — responsable installation et formation", [
            ("2019-2025", "Technicien de maintenance photovoltaïque, entreprise d'installation solaire, Douala et Edéa."),
            ("2016-2019", "Électricien bâtiment, chantiers résidentiels et commerciaux."),
        ], ["BTS électrotechnique (2016)", "Certificat d'installateur photovoltaïque (2020)"],
            "Installation solaire, maintenance, formation de techniciens, sécurité électrique.")),
        ("lettre_soutien_partenaire.pdf", support_letter(today)),
    ]


_DEMO_CACHE: dict = {}


def demo_bundle(today: date):
    """Call + pieces for the demonstration dossier, generated once per day (documents are deterministic for a given day)."""
    if _DEMO_CACHE.get("day") != today:
        _DEMO_CACHE.clear()
        _DEMO_CACHE.update(day=today, call=call_pdf(today), pieces=demo_pieces(today))
    return _DEMO_CACHE["call"], _DEMO_CACHE["pieces"]


def call_and_pieces(today: date):
    return call_pdf(today), demo_pieces(today), correction_pieces(today)
